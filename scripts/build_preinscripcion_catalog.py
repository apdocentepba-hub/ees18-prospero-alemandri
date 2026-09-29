#!/usr/bin/env python3
"""Build the 2027 preinscription primary-school snapshot from DGCyE official data."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DISTRICTS_PATH = ROOT / "data" / "preinscripcion-2027-distritos.json"
OUTPUT_PATH = ROOT / "data" / "preinscripcion-2027-primarias.json"
SOURCE_META_PATH = ROOT / "data" / "preinscripcion-2027-source.json"
SOURCE_URL = (
    "https://abc.gob.ar/secretarias/sites/default/files/2026-06/"
    "N%C3%B3mina%20de%20Establecimientos%20con%20Matr%C3%ADcula%20Inicial%202026.xlsx"
)
SOURCE_SHEET = "Nómina de Unidades de Servicio"


def norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-zA-Z0-9]+", "_", text).strip("_").lower()


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def sector_label(raw: object) -> str | None:
    value = norm(raw)
    if "estatal" in value or value == "oficial":
        return "Estatal"
    if "privad" in value:
        return "Privada"
    return None


def _header_map(ws) -> tuple[int, dict[str, int]]:
    required = {"codigo_distrito_dgcye", "nombre_del_establecimiento", "sector_de_gestion", "nivel"}
    for row_number, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 25), values_only=True), start=1):
        mapped = {norm(value): index for index, value in enumerate(row) if clean(value)}
        if required.issubset(mapped):
            return row_number, mapped
    raise RuntimeError("No se encontró la fila de encabezados esperada en la nómina oficial.")


def _value(row: tuple[object, ...], columns: dict[str, int], *names: str) -> str:
    for name in names:
        index = columns.get(name)
        if index is not None and index < len(row):
            value = clean(row[index])
            if value:
                return value
    return ""


def _district_code(raw: object, valid_codes: set[int]) -> int | None:
    digits = re.sub(r"\D", "", clean(raw))
    if not digits:
        return None
    code = int(digits)
    return code if code in valid_codes else None


def _school_number(clave: str) -> str:
    compact = re.sub(r"\W", "", clave or "")
    match = re.search(r"(\d{4})$", compact)
    return match.group(1) if match else ""


def build_rows_from_xlsx(source_path: Path, districts: list[dict[str, object]]) -> list[dict[str, object]]:
    district_names = {int(item["codigo"]): str(item["nombre"]) for item in districts}
    valid_codes = set(district_names)

    wb = load_workbook(source_path, read_only=True, data_only=True)
    if SOURCE_SHEET not in wb.sheetnames:
        raise RuntimeError(f"No se encontró la hoja oficial “{SOURCE_SHEET}”.")
    ws = wb[SOURCE_SHEET]
    header_row, columns = _header_map(ws)

    schools: dict[tuple[int, str], dict[str, object]] = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        nivel = _value(row, columns, "nivel")
        if "primar" not in norm(nivel):
            continue

        gestion = sector_label(_value(row, columns, "sector_de_gestion"))
        if not gestion:
            continue

        district_code = _district_code(_value(row, columns, "codigo_distrito_dgcye"), valid_codes)
        if district_code is None:
            continue

        clave = _value(row, columns, "clave_provincial")
        cue = _value(row, columns, "cue")
        anexo = _value(row, columns, "anexo")
        nombre = _value(row, columns, "nombre_del_establecimiento")
        localidad = _value(row, columns, "localidad")
        modalidad = _value(row, columns, "modalidad")
        codigo_organizacion = _value(row, columns, "codigo_de_organizacion")
        school_id = clave or (cue + ("-" + anexo if anexo else ""))

        if not school_id or not nombre:
            continue

        cue_anexo = cue + (("-" + anexo) if anexo else "") if cue else ""
        record = {
            "distritoCodigo": district_code,
            "distrito": district_names[district_code],
            "id": school_id,
            "nombre": nombre,
            "gestion": gestion,
            "clave": clave,
            "cueAnexo": cue_anexo,
            "nroEscuela": _school_number(clave),
            "modalidad": modalidad,
            "localidad": localidad,
            "codigoOrganizacion": codigo_organizacion,
        }
        schools.setdefault((district_code, school_id), record)

    result = list(schools.values())
    result.sort(key=lambda item: (int(item["distritoCodigo"]), norm(item["nombre"]), str(item["id"])))
    return result


def validate_snapshot(schools: list[dict[str, object]]) -> tuple[int, int]:
    if len(schools) < 1000:
        raise RuntimeError(f"El filtro devolvió solo {len(schools)} primarias; se aborta para no publicar un padrón incompleto.")

    avellaneda = [row for row in schools if row["distritoCodigo"] == 5]
    sectors = {row["gestion"] for row in avellaneda}
    if not {"Estatal", "Privada"}.issubset(sectors):
        raise RuntimeError("El padrón de Avellaneda no contiene gestión estatal y privada.")

    covered = {int(row["distritoCodigo"]) for row in schools}
    if len(covered) < 130:
        raise RuntimeError(f"El padrón cubre solo {len(covered)} distritos; se aborta.")
    return len(schools), len(covered)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-file", type=Path, required=True, help="XLSX oficial DGCyE ya descargado")
    args = parser.parse_args()

    districts = json.loads(DISTRICTS_PATH.read_text(encoding="utf-8"))
    if len(districts) != 137:
        raise RuntimeError(f"Se esperaban 137 distritos y hay {len(districts)}.")
    if not args.source_file.exists():
        raise RuntimeError(f"No existe el archivo fuente: {args.source_file}")

    schools = build_rows_from_xlsx(args.source_file, districts)
    school_count, district_count = validate_snapshot(schools)

    OUTPUT_PATH.write_text(json.dumps(schools, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    SOURCE_META_PATH.write_text(
        json.dumps(
            {
                "fuente": "DGCyE — Nómina de Establecimientos con Matrícula Inicial 2026",
                "url": SOURCE_URL,
                "hoja": SOURCE_SHEET,
                "relevamiento": "Relevamiento Inicial 2026",
                "generadoUtc": datetime.now(timezone.utc).isoformat(),
                "cantidadPrimarias": school_count,
                "distritosCubiertos": district_count,
                "filtro": "Nivel contiene PRIMAR; Sector de gestión Estatal o Privado/Privada",
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"Generadas {school_count} primarias en {district_count} distritos.")


if __name__ == "__main__":
    main()
